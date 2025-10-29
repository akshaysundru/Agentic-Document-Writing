from io import BytesIO
from docx import Document
from docx.shared import Pt
from flask import send_file
from bs4 import BeautifulSoup, Tag
import pymupdf4llm as pdf
import pathlib
from .constants import MARKDOWN_PATH, PDF_DIR
import markdown

def html_to_docx(html_content: str) -> BytesIO:
    """
    Converts simple HTML (paragraphs, <br>, tables) to a Word document in memory.
    """
    doc = Document()
    soup = BeautifulSoup(html_content, "html.parser")

    for element in soup.contents:
        if isinstance(element, Tag):
            if element.name == "p":
                p = doc.add_paragraph()
                p.add_run(element.get_text(separator="\n")).font.size = Pt(12)
            elif element.name == "br":
                doc.add_paragraph()  # add line break as a new paragraph
            elif element.name == "table":
                rows = element.find_all("tr")
                if rows:
                    table = doc.add_table(rows=0, cols=len(rows[0].find_all(["td", "th"])))
                    for row in rows:
                        cells = row.find_all(["td", "th"])
                        row_cells = table.add_row().cells
                        for i, cell in enumerate(cells):
                            row_cells[i].text = cell.get_text()
        else:
            # For plain strings or other non-tag content, add as a paragraph
            text = str(element).strip()
            if text:
                doc.add_paragraph(text)

    # Save to in-memory file
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

def pdf_to_markdown(directory):
    directory = pathlib.Path(directory)
    pathlib.Path(MARKDOWN_PATH).mkdir(exist_ok=True)

    converted_files = []

    # Loop over all PDF files in the directory
    for file_path in directory.glob("*.pdf"):
        md_text = pdf.to_markdown(str(file_path))  # convert PDF to markdown
        output_file = pathlib.Path(MARKDOWN_PATH) / f"{file_path.stem}.md"
        output_file.write_text(md_text, encoding="utf-8")
        converted_files.append(str(output_file))

    return converted_files

def markdown_to_sections(md_text: str):
    html_body = markdown.markdown(md_text, extensions=['tables', 'fenced_code'])
    soup = BeautifulSoup(html_body, "html.parser")

    sections = []
    current_section = None

    for node in soup.children:
        if isinstance(node, str):
            continue  # skip plain text
        if node.name and node.name.startswith('h') and node.name[1].isdigit(): # type: ignore
            # Save previous section
            if current_section:
                sections.append(current_section)
            # Start new section (include header in content)
            current_section = {"content": str(node)}
        else:
            if current_section:
                current_section["content"] += str(node)
            else:
                current_section = {"content": str(node)}

    if current_section:
        sections.append(current_section)

    return sections


if __name__ == "__main__":
    pdf_to_markdown(directory=PDF_DIR)
